"""Pack rendered slide PNGs into a PDF and a 16:9 PowerPoint (one full-bleed image per slide)."""
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Emu

NAME = "campus-dating-show-slides"


def main() -> None:
    png_dir, out_dir = Path(sys.argv[1]), Path(sys.argv[2])
    pngs = sorted(png_dir.glob("slide-*.png"))
    if not pngs:
        raise SystemExit("no slides rendered")

    pages = [Image.open(p).convert("RGB") for p in pngs]
    pdf_path = out_dir / f"{NAME}.pdf"
    pages[0].save(pdf_path, save_all=True, append_images=pages[1:], resolution=144)

    # The cotton texture makes PNGs heavy; high-quality JPEG keeps the PPTX shareable.
    jpgs = []
    for page, png in zip(pages, pngs):
        jpg = png.with_suffix(".jpg")
        page.save(jpg, quality=90, optimize=True)
        jpgs.append(jpg)

    deck = Presentation()
    deck.slide_width = Emu(12192000)
    deck.slide_height = Emu(6858000)
    blank = deck.slide_layouts[6]
    for jpg in jpgs:
        slide = deck.slides.add_slide(blank)
        slide.shapes.add_picture(str(jpg), 0, 0, width=deck.slide_width, height=deck.slide_height)
    pptx_path = out_dir / f"{NAME}.pptx"
    deck.save(pptx_path)
    print("wrote", pdf_path)
    print("wrote", pptx_path)


if __name__ == "__main__":
    main()
